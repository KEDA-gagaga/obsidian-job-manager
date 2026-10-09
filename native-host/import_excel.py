"""Read the committed recruitment workbook without modifying it."""
import argparse
import datetime as dt
import hashlib
import io
import json
import re
import unicodedata
from pathlib import Path

import openpyxl

DEFAULT_EXCEL = ""
ALIASES = {}


def key(text):
    return re.sub(r"[\s·•]+", "", unicodedata.normalize("NFKC", text)).casefold()


def company_name(value):
    text = str(value or "").strip()
    if not text or text in {"多家参会单位", "上海静安区", "武汉市", "临夏州", "南京江北新区", "四川省",
                           "湖南省", "贵州省", "陕西省", "广州荔湾", "香港人才办", "环境学院",
                           "安徽省省属国有企业", "国防科工重点单位", "国家某部委及其直属机构",
                           "广州开发区/广州市黄埔区事业单位", "中国机械工业集团、苏美达股份"}:
        return None
    if text in ALIASES:
        return ALIASES[text]
    if any(word in text for word in ("双选会", "招聘会", "就业指导", "参会单位", "人才招聘周")):
        return None
    return text


def value(cell):
    v = cell.value
    if isinstance(v, (dt.datetime, dt.date)):
        return v.isoformat(sep=" ") if isinstance(v, dt.datetime) else v.isoformat()
    return "" if v is None else v


def read_excel(path):
    p = Path(path)
    # Detect both an open Excel editor and replacement while reading.
    if p.with_name("~$" + p.name).exists():
        raise RuntimeError("Excel 正在编辑，关闭表格后再刷新。")
    before = p.stat()
    content = p.read_bytes()
    after = p.stat()
    if (before.st_mtime_ns, before.st_size) != (after.st_mtime_ns, after.st_size):
        raise RuntimeError("Excel 正在更新，请稍后再刷新。")
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    records = []
    companies = {}
    excluded = []
    for ws in wb:
        if ws.title not in {"招聘会信息", "投递信息"}:
            continue
        heads = [str(c.value or "") for c in ws[1]]
        if "企业" not in heads or "记录ID" not in heads:
            raise RuntimeError(f"{ws.title} 缺少企业或记录ID，未导入。")
        for row in list(ws.iter_rows())[1:]:
            if not any(c.value is not None for c in row):
                continue
            fields = {heads[i]: value(c) for i, c in enumerate(row) if heads[i]}
            links = {heads[i]: c.hyperlink.target for i, c in enumerate(row)
                     if c.hyperlink and c.hyperlink.target}
            rid = str(fields.get("记录ID") or "")
            if not rid:
                rid = "source-" + hashlib.sha256((ws.title + json.dumps(fields, ensure_ascii=False, sort_keys=True)).encode()).hexdigest()[:20]
            name = company_name(fields.get("企业"))
            cid = "company-" + hashlib.sha256(key(name).encode()).hexdigest()[:16] if name else None
            record = {"id": rid, "sheet": ws.title, "row": row[0].row, "companyId": cid,
                      "fields": fields, "links": links}
            records.append(record)
            if name:
                c = companies.setdefault(cid, {"id": cid, "name": name, "aliases": [], "recordIds": [],
                                              "索引": ["[[工作投递管理系统]]"]})
                original = str(fields.get("企业") or "").strip()
                if original not in c["aliases"]:
                    c["aliases"].append(original)
                c["recordIds"].append(rid)
            else:
                excluded.append(rid)
    ids = [r["id"] for r in records]
    if len(ids) != len(set(ids)):
        raise RuntimeError("原表包含重复记录ID，需先核对后导入。")
    initial_jobs = []
    for r in records:
        f = r["fields"]
        if r["companyId"] and f.get("我的投递状态") == "已投递":
            initial_jobs.append({"id": "job-" + r["id"], "companyId": r["companyId"],
                "name": f.get("招聘计划/岗位") or "已投递岗位（待补名称）", "sourceId": r["id"],
                "batch": f.get("招聘批次") or "", "location": f.get("工作地点") or "",
                "url": r["links"].get("投递入口", ""), "deadline": f.get("投递截止时间") or "",
                "stage": "已投递", "appliedDate": str(f.get("我的投递日期") or "")[:10],
                "resume": "", "notes": f.get("个人备注") or "", "createdAt": dt.date.today().isoformat(),
                "索引": ["[[工作投递管理系统]]"]})
    return {"sourcePath": str(p), "sourceHash": hashlib.sha256(content).hexdigest(),
            "readAt": dt.datetime.now(dt.timezone(dt.timedelta(hours=8))).isoformat(),
            "companies": sorted(companies.values(), key=lambda c: c["name"]), "records": records,
            "initialJobs": initial_jobs, "unassignedRecordIds": excluded,
            "counts": {"companies": len(companies), "records": len(records),
                       "delivery": sum(r["sheet"] == "投递信息" for r in records),
                       "events": sum(r["sheet"] == "招聘会信息" for r in records),
                       "unassigned": len(excluded)}}


if __name__ == "__main__":
    a = argparse.ArgumentParser()
    a.add_argument("--excel", required=True)
    args = a.parse_args()
    print(json.dumps(read_excel(args.excel), ensure_ascii=False))
